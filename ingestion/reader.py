"""Format-agnostic row reading for organizer CSV/XLSX exports.

Exports differ in everything but their content: comma, semicolon or tab between the cells (a
spreadsheet in a European locale writes semicolons), UTF-8, UTF-16 or Windows-1252 text, a title
and a blank line above the header, times as text or as spreadsheet time cells. All of that is
absorbed here, so the rest of ingestion sees one thing: a header row and rows of stripped strings.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import zipfile
from pathlib import Path

from defusedxml import ElementTree as SafeElementTree
from typing import Callable

from openpyxl import load_workbook


# A results file is a few thousand rows at most; a million-row sheet (or a zip bomb pretending to be
# one) is refused before it costs minutes of CPU.
MAX_ROWS = 50_000
# How far down the header row may sit under titles, logos and blank lines.
HEADER_SEARCH_ROWS = 30
# A results row has a dozen cells; timing exports with every split have a hundred. The cap is what
# stops a file that is wide instead of long: finding the header reads every cell of the first
# HEADER_SEARCH_ROWS rows and runs two regular expressions on each, so 60,000 columns cost seconds
# of a core per upload. The .xlsx reader has had this cap (MAX_XLSX_COLUMNS); CSV had none.
MAX_COLUMNS = 200
# No name, time or bib is longer than this. The cap is what stops a file whose cost is in the size
# of its cells rather than their number: two hundred headers of a hundred kilobytes each are a
# legal results file by every other measure, and they came back in the answer three times over.
MAX_CELL_CHARS = 512
# Rows times columns is what costs memory, and each limit on its own allows ten million cells:
# fifty thousand rows of two hundred columns is a 20 MB upload that cost 412 MB and four and a
# half seconds. Neither dimension is worth narrowing on its own -- fifty thousand finishers is a
# real race, and a timing export with every split is really that wide -- but no results file is
# both at once.
MAX_CELLS = 2_000_000
# Anything below a space except tab and newline. A results file has no business carrying them, and
# a NUL is not storable in a Postgres text column, so one in a name became an error at the end of
# a long upload rather than a word about the file.
_CONTROLS = {code: None for code in range(32) if code not in (9, 10, 13)} | {127: None}
# The characters that tell a browser to lay text out right to left. A name carrying one renders
# with the rank and the time around it in the wrong order, which is a leaderboard that lies
# without any string in it being wrong. No name needs them; Arabic and Hebrew read right to left
# on their own letters.
_CONTROLS |= {code: None for code in (0x200E, 0x200F, 0x061C, *range(0x202A, 0x202F), *range(0x2066, 0x206A))}
TEXT_SUFFIXES = (".csv", ".txt", ".tsv")
SHEET_SUFFIXES = (".xlsx", ".xlsm")

HeaderScore = Callable[[list[str]], int]


def read_rows(path: Path, header_score: HeaderScore | None = None) -> list[dict[str, str]]:
    """Read a CSV or XLSX file into a list of row dicts keyed by the original header text.

    All cell values are returned as stripped strings, regardless of the underlying spreadsheet
    cell type, so callers never need to branch on file format. `header_score` says how much a row
    looks like the header (the number of cells it recognises); with it, title lines above the
    header are skipped. Without it the first row is the header.
    """
    headers, rows = read_table(path, header_score)
    return [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in rows]


def read_table(path: Path, header_score: HeaderScore | None = None) -> tuple[list[str], list[list[str]]]:
    """The header cells and the data rows (lists of stripped strings, blank rows dropped)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        table = _read_text(path)
    elif suffix in SHEET_SUFFIXES:
        table = _read_xlsx(path)
    elif suffix == ".xls":
        raise ValueError("this is the old Excel format (.xls): open it and save it as .xlsx or CSV, then upload that")
    else:
        raise ValueError(f"Unsupported file type: {suffix!r} (upload a .csv or .xlsx file)")

    table = [row for row in table if any(cell for cell in row)]
    if not table:
        return [], []
    start = _header_index(table, header_score)
    headers = _unique(table[start])
    rows = table[start + 1 :]
    if len(rows) > MAX_ROWS:
        raise ValueError(f"the file has more than {MAX_ROWS} rows; split it per race distance")
    return headers, rows


def _header_index(table: list[list[str]], header_score: HeaderScore | None) -> int:
    if header_score is None:
        return 0
    best, best_score = 0, 0
    for index, row in enumerate(table[:HEADER_SEARCH_ROWS]):
        score = header_score(row)
        if score > best_score:  # the first of equally good rows: the header comes before the data
            best, best_score = index, score
    return best if best_score >= 2 else 0


def _unique(headers: list[str]) -> list[str]:
    """Header texts, with a repeated one numbered ("Time", "Time (2)") so no column hides another."""
    seen: dict[str, int] = {}
    result = []
    for header in headers:
        count = seen.get(header, 0) + 1
        seen[header] = count
        result.append(header if count == 1 or not header else f"{header} ({count})")
    return result


# ---------------------------------------------------------------------------- text files


def _cell(value: str | None) -> str:
    """One cell as it is kept: no control characters, and no longer than a name can be."""
    return (value or "").translate(_CONTROLS).strip()[:MAX_CELL_CHARS]


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    # "Unicode text" from a spreadsheet, without a byte-order mark: every other byte is zero.
    # Which other one says which way round it is, and only the little-endian half used to be read.
    if len(data) >= 4:
        for first, encoding in ((1, "utf-16-le"), (0, "utf-16-be")):
            if data[first : first + 1] == b"\x00" and data[first + 2 : first + 3] == b"\x00":
                try:
                    return data.decode(encoding)
                except UnicodeDecodeError:
                    pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")  # what Excel on Windows writes for "CSV"


def _delimiter(text: str) -> str:
    """The cell separator: the candidate that splits the first lines into the most cells, the same
    number of times on most lines. Quoted cells are respected."""
    lines = [line for line in text.splitlines()[:40] if line.strip()]
    best, best_score = ",", 0.0
    for candidate in (",", ";", "\t", "|"):
        try:
            counts = [len(row) for row in csv.reader(lines, delimiter=candidate)]
        except csv.Error:
            # A line with no such separator is one enormous field, and the csv module refuses a
            # field over 128 KB. That says this is not the separator, not that the file is broken.
            continue
        wide = [count for count in counts if count > 1]
        if not wide:
            continue
        typical = max(set(wide), key=wide.count)
        score = wide.count(typical) * typical
        if score > best_score:
            best, best_score = candidate, score
    return best


def _read_text(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    delimiter = "\t" if path.suffix.lower() == ".tsv" else _delimiter(text)
    table: list[list[str]] = []
    cells = 0
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            if len(table) > MAX_ROWS + HEADER_SEARCH_ROWS:
                raise ValueError(f"the file has more than {MAX_ROWS} rows; split it per race distance")
            # Past MAX_COLUMNS the rest of the row is dropped, as the .xlsx reader drops it.
            kept = [_cell(cell) for cell in row[:MAX_COLUMNS]]
            cells += len(kept)
            if cells > MAX_CELLS:
                raise ValueError(_TOO_MANY_CELLS)
            table.append(kept)
    except csv.Error as error:
        # csv.Error is not a ValueError, so it used to leave the reader as an unhandled error and
        # the upload answered 500 instead of saying which file could not be read.
        raise ValueError("this file could not be read as a table: save it again as CSV or .xlsx and upload that") from error
    return table


# ---------------------------------------------------------------------------- spreadsheets


# An .xlsx is a zip. A few kilobytes can unpack to gigabytes (a "zip bomb"), so what the archive
# claims to hold is checked before anything is unpacked. A real results sheet of 50,000 rows is a
# few megabytes unpacked; the old 200 MB allowance was far more than any of them needs.
MAX_XLSX_UNPACKED_BYTES = 25_000_000
# The one part that must be checked on its own. openpyxl reads xl/sharedStrings.xml into a Python
# list in full before a single row is read, so read_only=True, reset_dimensions(), MAX_ROWS and
# MAX_XLSX_COLUMNS all come too late to help. A 0.12 MB upload declaring 50 MB of shared strings
# cost 59 seconds of a core and 367 MB; the row and column caps never saw it.
MAX_XLSX_SHARED_STRINGS_BYTES = 6_000_000
_SHARED_STRINGS = "xl/sharedstrings.xml"
# openpyxl does not look for that name. It reads [Content_Types].xml and takes whatever part is
# declared with this content type, so the strings can sit at any path in the archive and a check
# on the usual name alone is bypassed by renaming the part.
_SHARED_STRINGS_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedstrings+xml"
MAX_XLSX_ENTRIES = 2_000
MAX_XLSX_COLUMNS = 200  # a results sheet has a dozen; timing exports with every split, a hundred
_TOO_MANY_CELLS = (
    "this file holds more cells than a results sheet can: keep the columns the scorer reads "
    "(a rank, a time, a name, a gender) and split it per race distance"
)
_NOT_A_WORKBOOK = "this file is named .xlsx but is not an Excel workbook (it may be a CSV that was renamed, or a damaged download): open it in a spreadsheet, save it as .xlsx or CSV, and upload that"


def _refuse_oversized_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, OSError) as error:
        raise ValueError(_NOT_A_WORKBOOK) from error
    too_big = "this workbook is far larger inside than a results sheet can be; save the results as CSV and upload that"
    if len(entries) > MAX_XLSX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_XLSX_UNPACKED_BYTES:
        raise ValueError(too_big)
    names = _shared_strings_parts(path) | {_SHARED_STRINGS}
    shared = sum(entry.file_size for entry in entries if entry.filename.lower().lstrip("/") in names)
    if shared > MAX_XLSX_SHARED_STRINGS_BYTES:
        raise ValueError(too_big)


def _shared_strings_parts(path: Path) -> set[str]:
    """Which parts of this archive the workbook declares as its shared strings, read the way
    openpyxl reads them: from [Content_Types].xml, by content type rather than by name."""
    try:
        with zipfile.ZipFile(path) as archive:
            with archive.open("[Content_Types].xml") as handle:
                raw = handle.read(1_000_000)
        root = SafeElementTree.fromstring(raw)
    except Exception:  # noqa: BLE001 - no content types, unreadable, or not XML: the name check stands
        return set()
    found = set()
    for override in root.iter():
        if not override.tag.endswith("Override"):
            continue
        if (override.get("ContentType") or "").strip().lower() == _SHARED_STRINGS_TYPE:
            found.add((override.get("PartName") or "").strip().lower().lstrip("/"))
    return {name for name in found if name}


def _read_xlsx(path: Path) -> list[list[str]]:
    _refuse_oversized_archive(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - openpyxl raises a zoo of errors for non-spreadsheets
        raise ValueError("this file is named .xlsx but is not an Excel workbook (it may be a CSV that was renamed, or a damaged download): open it in a spreadsheet, save it as .xlsx or CSV, and upload that") from error
    rows = None
    try:
        sheet = workbook.active
        if sheet is None:
            return []
        # A sheet says how large it is, and a reader that believes it pads every row to that width
        # and fills every gap between two rows with empty ones. A workbook of 1.5 KB that claimed
        # 16,384 columns and put one cell in row 60,000 cost 14 seconds of a core, and the public
        # scorer takes workbooks from anyone. So the sheet's claim is dropped, no row is read past
        # MAX_XLSX_COLUMNS, and every row counts towards the limit, the empty ones too.
        sheet.reset_dimensions()
        limit = MAX_ROWS + HEADER_SEARCH_ROWS
        rows = sheet.iter_rows(values_only=True, max_row=limit + 1, max_col=MAX_XLSX_COLUMNS)
        table: list[list[str]] = []
        cells = 0
        for seen, raw_row in enumerate(rows, start=1):
            if seen > limit:
                raise ValueError(f"the file has more than {MAX_ROWS} rows; split it per race distance")
            if all(cell is None for cell in raw_row):
                continue
            kept = [_cell(_stringify(cell)) for cell in raw_row]
            # A sheet is read to a fixed width, so a three-column sheet comes back padded out to
            # two hundred. The padding is not data and must not count against the cell budget, nor
            # be stored fifty thousand times over.
            while kept and not kept[-1]:
                kept.pop()
            cells += len(kept)
            if cells > MAX_CELLS:
                raise ValueError(_TOO_MANY_CELLS)
            table.append(kept)
        return table
    finally:
        # Read-only workbooks stream from the zip: close the row generator (its open stream) and the
        # archive, or the caller cannot delete the temp file on Windows.
        try:
            if rows is not None:
                rows.close()
        except Exception:  # noqa: BLE001
            pass
        workbook.close()


def _stringify(cell: object) -> str:
    """Normalize a spreadsheet cell (which may be a native date/time) to plain text."""
    if cell is None:
        return ""
    if isinstance(cell, dt.timedelta):  # a duration cell, e.g. a finish time beyond 24 hours
        seconds = round(cell.total_seconds())
        return f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"
    if isinstance(cell, dt.datetime):
        # A time beyond 24 h typed into a time cell comes back as a date in January 1900.
        if cell.year <= 1900:
            days = (cell.date() - dt.date(1899, 12, 31)).days if cell.year == 1900 else 0
            seconds = max(days, 0) * 86400 + cell.hour * 3600 + cell.minute * 60 + cell.second
            return f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"
        return cell.date().isoformat() if cell.time() == dt.time(0, 0) else cell.isoformat(sep=" ")
    if isinstance(cell, dt.date):
        return cell.isoformat()
    if isinstance(cell, dt.time):
        return cell.strftime("%H:%M:%S")
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))  # a bib or a rank kept as a number: "12", not "12.0"
    return str(cell).strip()
